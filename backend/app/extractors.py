import csv
import io
import json
from pathlib import Path
from typing import List

import pymupdf
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_core.documents import Document
from openpyxl import load_workbook
from pptx import Presentation


PLAIN_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".json", ".jsonl", ".xml", ".yaml", ".yml",
    ".py", ".js", ".jsx", ".ts", ".tsx", ".css", ".scss", ".sql", ".sh",
    ".java", ".go", ".rs", ".c", ".h", ".cpp", ".toml", ".ini", ".log",
}
SUPPORTED_EXTENSIONS = PLAIN_EXTENSIONS | {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".html", ".htm"}


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_documents(filename: str, data: bytes) -> List[Document]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{extension or 'unknown'}'. Supported: {supported}")

    if extension == ".pdf":
        return _extract_pdf(filename, data)
    if extension == ".docx":
        return _extract_docx(filename, data)
    if extension == ".pptx":
        return _extract_pptx(filename, data)
    if extension == ".xlsx":
        return _extract_xlsx(filename, data)
    if extension == ".csv":
        return _extract_csv(filename, data)
    if extension in {".html", ".htm"}:
        soup = BeautifulSoup(_decode(data), "html.parser")
        return [Document(page_content=soup.get_text("\n", strip=True), metadata={"source": filename})]

    text = _decode(data)
    if extension == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except json.JSONDecodeError:
            pass
    return [Document(page_content=text, metadata={"source": filename})]


def _extract_pdf(filename: str, data: bytes) -> List[Document]:
    output: List[Document] = []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text", sort=True).strip()
            if text:
                output.append(Document(page_content=text, metadata={"source": filename, "page": page_number}))
    if not output:
        raise ValueError("No text was found. This may be an image-only PDF that requires OCR.")
    return output


def _extract_docx(filename: str, data: bytes) -> List[Document]:
    doc = DocxDocument(io.BytesIO(data))
    blocks = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))
    return [Document(page_content="\n\n".join(blocks), metadata={"source": filename})]


def _extract_pptx(filename: str, data: bytes) -> List[Document]:
    presentation = Presentation(io.BytesIO(data))
    output: List[Document] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )
        if text:
            output.append(Document(page_content=text, metadata={"source": filename, "page": slide_number, "section": "slide"}))
    return output


def _extract_xlsx(filename: str, data: bytes) -> List[Document]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    output: List[Document] = []
    for sheet in workbook.worksheets:
        lines = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                lines.append(" | ".join(values))
        if lines:
            output.append(Document(page_content="\n".join(lines), metadata={"source": filename, "section": sheet.title}))
    return output


def _extract_csv(filename: str, data: bytes) -> List[Document]:
    text = _decode(data)
    rows = list(csv.reader(io.StringIO(text)))
    normalized = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    return [Document(page_content=normalized, metadata={"source": filename})]

